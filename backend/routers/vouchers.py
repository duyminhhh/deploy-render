"""
Voucher Router - quản lý phiếu nhập/xuất kho với đầy đủ workflow
"""
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
from typing import List, Optional
from datetime import datetime, timezone
import io, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from database import get_db
import models, schemas
from auth import get_current_user, get_current_user_from_token, require_role

router = APIRouter(prefix="/vouchers", tags=["Vouchers"])
manager_only = require_role("MANAGER")
staff_or_above = require_role("MANAGER", "STAFF")

INVESTIGATION_THRESHOLD = 0.1   # chênh lệch > 10% → INVESTIGATING

def _now():
    return datetime.now(timezone.utc)

def _gen_code(db: Session, type_: str) -> str:
    year = datetime.utcnow().year
    prefix = "NK" if type_ == "IN" else "XK"
    count = db.query(models.Voucher).filter(
        models.Voucher.code.like(f"{prefix}-{year}-%")
    ).count() + 1
    return f"{prefix}-{year}-{count:04d}"

def _log(db: Session, voucher_id: int, action: str, from_status, to_status, user_id: int, note: str = None):
    db.add(models.VoucherLog(
        voucher_id=voucher_id,
        action=action,
        from_status=from_status.value if from_status else None,
        to_status=to_status.value if hasattr(to_status, "value") else str(to_status),
        by_user_id=user_id,
        note=note
    ))

def _load_voucher(db, voucher_id):
    return db.query(models.Voucher).options(
        joinedload(models.Voucher.items).joinedload(models.VoucherItem.product),
        joinedload(models.Voucher.creator),
        joinedload(models.Voucher.approver),
        joinedload(models.Voucher.completer),
    ).filter(models.Voucher.id == voucher_id).first()

# ─── CREATE (Manager) → DRAFT ─────────────────────────────────────────────────
@router.post("/", response_model=schemas.VoucherOut, status_code=201)
def create_voucher(
    data: schemas.VoucherCreate,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(manager_only)
):
    voucher = models.Voucher(
        type=data.type,
        title=data.title,
        note=data.note,
        valid_until=data.valid_until,
        created_by=current_user.id,
        status=models.VoucherStatusEnum.DRAFT
    )
    db.add(voucher)
    db.flush()
    for item in data.items:
        product = db.query(models.Product).filter(models.Product.id == item.product_id).first()
        if not product:
            raise HTTPException(status_code=404, detail=f"Product {item.product_id} not found")
        db.add(models.VoucherItem(
            voucher_id=voucher.id,
            product_id=item.product_id,
            approved_quantity=item.approved_quantity
        ))
    _log(db, voucher.id, "CREATED", None, models.VoucherStatusEnum.DRAFT, current_user.id)
    db.commit()
    return _load_voucher(db, voucher.id)

# ─── UPDATE DRAFT (Manager only) ─────────────────────────────────────────────
@router.put("/{voucher_id}", response_model=schemas.VoucherOut)
def update_voucher(
    voucher_id: int,
    data: schemas.VoucherUpdate,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(manager_only)
):
    v = db.query(models.Voucher).filter(models.Voucher.id == voucher_id).first()
    if not v:
        raise HTTPException(404, "Voucher not found")
    if v.status not in (models.VoucherStatusEnum.DRAFT, models.VoucherStatusEnum.PENDING):
        raise HTTPException(400, "Chỉ có thể sửa phiếu ở trạng thái DRAFT hoặc PENDING")
    if data.title is not None: v.title = data.title
    if data.note is not None: v.note = data.note
    if data.valid_until is not None: v.valid_until = data.valid_until
    if data.items is not None:
        # remove old items, re-add
        db.query(models.VoucherItem).filter(models.VoucherItem.voucher_id == v.id).delete()
        for item in data.items:
            db.add(models.VoucherItem(
                voucher_id=v.id, product_id=item.product_id,
                approved_quantity=item.approved_quantity
            ))
    db.commit()
    return _load_voucher(db, voucher_id)

# ─── SUBMIT → PENDING (Manager) ───────────────────────────────────────────────
@router.post("/{voucher_id}/submit", response_model=schemas.VoucherOut)
def submit_voucher(
    voucher_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(manager_only)
):
    v = db.query(models.Voucher).filter(models.Voucher.id == voucher_id).first()
    if not v: raise HTTPException(404, "Voucher not found")
    if v.status != models.VoucherStatusEnum.DRAFT:
        raise HTTPException(400, "Chỉ phiếu DRAFT mới được submit")
    old = v.status
    v.status = models.VoucherStatusEnum.PENDING
    _log(db, v.id, "SUBMITTED", old, v.status, current_user.id)
    db.commit()
    return _load_voucher(db, voucher_id)

# ─── APPROVE → APPROVED (Manager) - sinh mã, khoá phiếu ─────────────────────
@router.post("/{voucher_id}/approve", response_model=schemas.VoucherOut)
def approve_voucher(
    voucher_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(manager_only)
):
    v = db.query(models.Voucher).filter(models.Voucher.id == voucher_id).first()
    if not v: raise HTTPException(404, "Voucher not found")
    if v.status not in (models.VoucherStatusEnum.PENDING, models.VoucherStatusEnum.DRAFT):
        raise HTTPException(400, f"Không thể duyệt phiếu ở trạng thái {v.status}")
    old = v.status
    v.status = models.VoucherStatusEnum.APPROVED
    v.code = _gen_code(db, v.type)
    v.approved_by = current_user.id
    v.approved_at = _now()
    _log(db, v.id, "APPROVED", old, v.status, current_user.id)
    db.commit()
    return _load_voucher(db, voucher_id)

# ─── REJECT (Manager) ─────────────────────────────────────────────────────────
@router.post("/{voucher_id}/reject", response_model=schemas.VoucherOut)
def reject_voucher(
    voucher_id: int,
    data: schemas.RejectRequest,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(manager_only)
):
    v = db.query(models.Voucher).filter(models.Voucher.id == voucher_id).first()
    if not v: raise HTTPException(404, "Voucher not found")
    if v.status not in (models.VoucherStatusEnum.DRAFT, models.VoucherStatusEnum.PENDING,
                        models.VoucherStatusEnum.APPROVED):
        raise HTTPException(400, f"Không thể từ chối phiếu {v.status}")
    old = v.status
    v.status = models.VoucherStatusEnum.REJECTED
    v.reject_reason = data.reason
    _log(db, v.id, "REJECTED", old, v.status, current_user.id, data.reason)
    db.commit()
    return _load_voucher(db, voucher_id)

# ─── CANCEL (Manager) ─────────────────────────────────────────────────────────
@router.post("/{voucher_id}/cancel", response_model=schemas.VoucherOut)
def cancel_voucher(
    voucher_id: int,
    data: schemas.RejectRequest,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(manager_only)
):
    v = db.query(models.Voucher).filter(models.Voucher.id == voucher_id).first()
    if not v: raise HTTPException(404, "Voucher not found")
    if v.status in (models.VoucherStatusEnum.COMPLETED, models.VoucherStatusEnum.CANCELLED):
        raise HTTPException(400, f"Không thể huỷ phiếu {v.status}")
    old = v.status
    v.status = models.VoucherStatusEnum.CANCELLED
    v.reject_reason = data.reason
    _log(db, v.id, "CANCELLED", old, v.status, current_user.id, data.reason)
    db.commit()
    return _load_voucher(db, voucher_id)

# ─── START (Staff nhận xử lý) → IN_PROGRESS ──────────────────────────────────
@router.post("/{voucher_id}/start", response_model=schemas.VoucherOut)
def start_voucher(
    voucher_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(staff_or_above)
):
    v = db.query(models.Voucher).filter(models.Voucher.id == voucher_id).first()
    if not v: raise HTTPException(404, "Voucher not found")
    if v.status != models.VoucherStatusEnum.APPROVED:
        raise HTTPException(400, "Chỉ phiếu APPROVED mới được bắt đầu xử lý")
    old = v.status
    v.status = models.VoucherStatusEnum.IN_PROGRESS
    _log(db, v.id, "STARTED", old, v.status, current_user.id)
    db.commit()
    return _load_voucher(db, voucher_id)

# ─── CONFIRM (Manager xác nhận) → COMPLETED / INVESTIGATING ──────────────────
@router.post("/{voucher_id}/confirm", response_model=schemas.VoucherOut)
def confirm_voucher(
    voucher_id: int,
    data: schemas.ConfirmRequest,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(manager_only)
):
    v = db.query(models.Voucher).filter(models.Voucher.id == voucher_id).first()
    if not v: raise HTTPException(404, "Voucher not found")
    if v.status != models.VoucherStatusEnum.IN_PROGRESS:
        raise HTTPException(400, "Chỉ phiếu IN_PROGRESS mới được xác nhận hoàn tất")

    # Reload items with product
    items = db.query(models.VoucherItem).options(
        joinedload(models.VoucherItem.product)
    ).filter(models.VoucherItem.voucher_id == v.id).all()

    # Check if all items have actual qty
    missing = [i for i in items if i.actual_quantity is None]
    if missing:
        raise HTTPException(400, f"Còn {len(missing)} sản phẩm chưa nhập số lượng thực tế")

    # Detect large discrepancy
    has_large_diff = False
    for item in items:
        if item.approved_quantity > 0:
            diff_ratio = abs(item.actual_quantity - item.approved_quantity) / item.approved_quantity
            if diff_ratio > INVESTIGATION_THRESHOLD:
                has_large_diff = True
                break

    old = v.status
    if has_large_diff:
        v.status = models.VoucherStatusEnum.INVESTIGATING
        action = "INVESTIGATING"
    else:
        v.status = models.VoucherStatusEnum.COMPLETED
        action = "COMPLETED"
        v.completed_by = current_user.id
        v.completed_at = _now()
        # Cập nhật tồn kho
        for item in items:
            product = item.product
            actual_qty = item.actual_quantity
            if v.type == models.TransactionTypeEnum.IN:
                product.quantity += actual_qty
                product.import_date = _now()
            else:
                if product.quantity < actual_qty:
                    raise HTTPException(400, f"Tồn kho không đủ cho '{product.name}': có {product.quantity}, cần {actual_qty}")
                product.quantity -= actual_qty
            # Ghi transaction log
            db.add(models.InventoryTransaction(
                product_id=item.product_id,
                type=v.type,
                quantity=actual_qty,
                note=data.note or f"Hoàn tất phiếu {v.code}",
                performed_by=current_user.id,
                voucher_id=v.id
            ))

    _log(db, v.id, action, old, v.status, current_user.id, data.note)
    db.commit()
    return _load_voucher(db, voucher_id)

# ─── FORCE COMPLETE from INVESTIGATING (Manager) ──────────────────────────────
@router.post("/{voucher_id}/force-complete", response_model=schemas.VoucherOut)
def force_complete_voucher(
    voucher_id: int,
    data: schemas.ConfirmRequest,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(manager_only)
):
    v = db.query(models.Voucher).filter(models.Voucher.id == voucher_id).first()
    if not v: raise HTTPException(404, "Voucher not found")
    if v.status != models.VoucherStatusEnum.INVESTIGATING:
        raise HTTPException(400, "Chỉ phiếu INVESTIGATING mới được force-complete")

    items = db.query(models.VoucherItem).options(
        joinedload(models.VoucherItem.product)
    ).filter(models.VoucherItem.voucher_id == v.id).all()

    old = v.status
    v.status = models.VoucherStatusEnum.COMPLETED
    v.completed_by = current_user.id
    v.completed_at = _now()

    for item in items:
        product = item.product
        actual_qty = item.actual_quantity or 0
        if v.type == models.TransactionTypeEnum.IN:
            product.quantity += actual_qty
            product.import_date = _now()
        else:
            product.quantity = max(0, product.quantity - actual_qty)
        db.add(models.InventoryTransaction(
            product_id=item.product_id, type=v.type, quantity=actual_qty,
            note=data.note or f"Force-complete phiếu {v.code} (sau điều tra)",
            performed_by=current_user.id, voucher_id=v.id
        ))

    _log(db, v.id, "FORCE_COMPLETED", old, v.status, current_user.id, data.note)
    db.commit()
    return _load_voucher(db, voucher_id)

# ─── LIST ─────────────────────────────────────────────────────────────────────
@router.get("/", response_model=List[schemas.VoucherOut])
def list_vouchers(
    type: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    q = db.query(models.Voucher).options(
        joinedload(models.Voucher.items).joinedload(models.VoucherItem.product),
        joinedload(models.Voucher.creator),
        joinedload(models.Voucher.approver),
        joinedload(models.Voucher.completer),
    )
    if type:
        q = q.filter(models.Voucher.type == type)
    if status:
        q = q.filter(models.Voucher.status == status)
    # STAFF chỉ thấy phiếu từ APPROVED trở đi
    if current_user.role == models.RoleEnum.STAFF:
        visible = [
            models.VoucherStatusEnum.APPROVED,
            models.VoucherStatusEnum.IN_PROGRESS,
            models.VoucherStatusEnum.COMPLETED,
            models.VoucherStatusEnum.INVESTIGATING,
        ]
        q = q.filter(models.Voucher.status.in_(visible))
    return q.order_by(models.Voucher.created_at.desc()).all()

# ─── GET ONE ──────────────────────────────────────────────────────────────────
@router.get("/{voucher_id}", response_model=schemas.VoucherOut)
def get_voucher(
    voucher_id: int,
    db: Session = Depends(get_db),
    _: models.User = Depends(get_current_user)
):
    v = _load_voucher(db, voucher_id)
    if not v: raise HTTPException(404, "Voucher not found")
    return v

# ─── LOGS ─────────────────────────────────────────────────────────────────────
@router.get("/{voucher_id}/logs", response_model=List[schemas.VoucherLogOut])
def get_voucher_logs(
    voucher_id: int,
    db: Session = Depends(get_db),
    _: models.User = Depends(get_current_user)
):
    return db.query(models.VoucherLog).options(
        joinedload(models.VoucherLog.by_user)
    ).filter(models.VoucherLog.voucher_id == voucher_id).order_by(models.VoucherLog.created_at).all()

# ─── PRINT Excel ──────────────────────────────────────────────────────────────
@router.get("/{voucher_id}/print")
def print_voucher(
    voucher_id: int,
    token: str = Query(..., description="JWT token"),
    db: Session = Depends(get_db),
):
    """Xuất phiếu kho ra file Excel — xác thực qua query param ?token=..."""
    # Xác thực token từ query param (vì browser download không gửi header)
    get_current_user_from_token(token, db)

    v = _load_voucher(db, voucher_id)
    if not v:
        raise HTTPException(404, "Voucher not found")

    STATUS_VI = {
        "DRAFT": "Nháp", "APPROVED": "Đã duyệt", "IN_PROGRESS": "Đang xử lý",
        "COMPLETED": "Hoàn tất", "REJECTED": "Từ chối",
        "CANCELLED": "Đã huỷ", "INVESTIGATING": "Đang điều tra"
    }
    TYPE_VI = {"IN": "NHẬP KHO", "OUT": "XUẤT KHO"}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Phiếu kho"

    # ── Styles ─────────────────────────────────────────────────────────────
    def font(bold=False, size=11, color="000000", italic=False):
        return Font(name="Arial", bold=bold, size=size, color=color, italic=italic)

    def fill(hex_color):
        return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")

    def border_all():
        s = Side(style="thin", color="BFBFBF")
        return Border(left=s, right=s, top=s, bottom=s)

    def align(h="left", v="center", wrap=False):
        return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

    # Column widths
    for col, w in zip("ABCDEFGH", [6, 32, 10, 14, 14, 14, 18, 20]):
        ws.column_dimensions[get_column_letter(col_num)].width = w             if (col_num := ord(col) - ord('A') + 1) else w

    row = 1

    # ── Header banner ───────────────────────────────────────────────────────
    ws.merge_cells(f"A{row}:H{row}")
    c = ws[f"A{row}"]
    c.value = f"PHIẾU {TYPE_VI.get(v.type, v.type)}"
    c.font = Font(name="Arial", bold=True, size=18, color="FFFFFF")
    c.fill = fill("1E3A8A")
    c.alignment = align("center")
    ws.row_dimensions[row].height = 44
    row += 1

    ws.merge_cells(f"A{row}:H{row}")
    c = ws[f"A{row}"]
    c.value = "BÁch HOÁ ĐỎ – Hệ thống quản lý kho BHD Storage"
    c.font = font(italic=True, size=10, color="6B7280")
    c.alignment = align("center")
    ws.row_dimensions[row].height = 18
    row += 2

    # ── Thông tin phiếu (2 cột) ─────────────────────────────────────────────
    def info_row(ws, row, label1, val1, label2, val2):
        ws[f"A{row}"].value = label1
        ws[f"A{row}"].font = font(bold=True, size=11)
        ws[f"A{row}"].fill = fill("F1F5F9")
        ws.merge_cells(f"B{row}:C{row}")
        ws[f"B{row}"].value = val1
        ws[f"B{row}"].font = font(size=11)

        ws[f"E{row}"].value = label2
        ws[f"E{row}"].font = font(bold=True, size=11)
        ws[f"E{row}"].fill = fill("F1F5F9")
        ws.merge_cells(f"F{row}:H{row}")
        ws[f"F{row}"].value = val2
        ws[f"F{row}"].font = font(size=11)

        for col in "ABCEFGH":
            ws[f"{col}{row}"].border = border_all()
            ws[f"{col}{row}"].alignment = align()
        ws.row_dimensions[row].height = 22

    info_row(ws, row, "Mã phiếu:", v.code or "(chưa có)", "Loại phiếu:", TYPE_VI.get(v.type, v.type)); row += 1
    info_row(ws, row, "Tiêu đề:", v.title or "", "Trạng thái:", STATUS_VI.get(v.status, v.status)); row += 1
    info_row(ws, row, "Người tạo:", v.creator.full_name or v.creator.username if v.creator else "–",
             "Ngày tạo:", v.created_at.strftime("%d/%m/%Y %H:%M") if v.created_at else "–"); row += 1
    info_row(ws, row, "Người duyệt:", v.approver.full_name or v.approver.username if v.approver else "–",
             "Ngày duyệt:", v.approved_at.strftime("%d/%m/%Y %H:%M") if v.approved_at else "–"); row += 1
    info_row(ws, row, "Ghi chú:", v.note or "", "Hạn hiệu lực:",
             v.valid_until.strftime("%d/%m/%Y") if v.valid_until else "Không giới hạn"); row += 2

    # ── Tiêu đề bảng hàng hoá ───────────────────────────────────────────────
    headers = ["STT", "Tên sản phẩm", "ĐV", "SL phê duyệt", "SL thực tế", "Chênh lệch", "Đơn giá (đ)", "Thành tiền (đ)"]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=ci, value=h)
        c.font = font(bold=True, size=11, color="FFFFFF")
        c.fill = fill("1D4ED8")
        c.alignment = align("center")
        c.border = border_all()
    ws.row_dimensions[row].height = 26
    row += 1

    # ── Dữ liệu hàng hoá ────────────────────────────────────────────────────
    total_approved = total_actual = total_value = 0
    for idx, item in enumerate(v.items, 1):
        actual = item.actual_quantity
        diff = (actual - item.approved_quantity) if actual is not None else None
        price = item.product.price if item.product else 0
        value = (actual or item.approved_quantity) * price

        row_bg = "EFF6FF" if idx % 2 == 0 else "FFFFFF"
        diff_color = "DC2626" if diff is not None and diff < 0 else                      "16A34A" if diff is not None and diff > 0 else "374151"

        vals = [
            idx,
            item.product.name if item.product else f"#{item.product_id}",
            item.product.unit if item.product else "",
            item.approved_quantity,
            actual if actual is not None else "–",
            (f"+{diff}" if diff and diff > 0 else str(diff)) if diff is not None else "–",
            f"{price:,.0f}" if price else "–",
            f"{value:,.0f}" if price else "–",
        ]
        for ci, val in enumerate(vals, 1):
            c = ws.cell(row=row, column=ci, value=val)
            c.border = border_all()
            c.alignment = align("center" if ci != 2 else "left")
            c.font = font(
                bold=(ci == 6 and diff is not None and diff < 0),
                color=diff_color if ci == 6 else "000000"
            )
            c.fill = fill(row_bg)
        ws.row_dimensions[row].height = 20

        total_approved += item.approved_quantity
        total_actual   += actual or 0
        total_value    += value
        row += 1

    # ── Dòng tổng ─────────────────────────────────────────────────────────
    ws.merge_cells(f"A{row}:C{row}")
    ws[f"A{row}"].value = "TỔNG CỘNG"
    ws[f"A{row}"].font = font(bold=True, size=11, color="FFFFFF")
    ws[f"A{row}"].fill = fill("1E3A8A")
    ws[f"A{row}"].alignment = align("center")

    for ci, val in enumerate([None, None, None, total_approved,
                               total_actual or "–", None, None, f"{total_value:,.0f}"], 1):
        if val is None: continue
        c = ws.cell(row=row, column=ci, value=val)
        c.font = font(bold=True, size=11, color="FFFFFF")
        c.fill = fill("1E3A8A")
        c.alignment = align("center")
        c.border = border_all()
    for ci in range(1, 9):
        ws.cell(row=row, column=ci).border = border_all()
    ws.row_dimensions[row].height = 24
    row += 3

    # ── Ký tên ─────────────────────────────────────────────────────────────
    ws.merge_cells(f"A{row}:C{row}")
    ws[f"A{row}"].value = "NGƯỜI PHÊ DUYỆT"
    ws[f"A{row}"].font = font(bold=True, size=11)
    ws[f"A{row}"].alignment = align("center")

    ws.merge_cells(f"F{row}:H{row}")
    ws[f"F{row}"].value = "NGƯỜI THỰC HIỆN"
    ws[f"F{row}"].font = font(bold=True, size=11)
    ws[f"F{row}"].alignment = align("center")
    ws.row_dimensions[row].height = 20
    row += 4

    ws.merge_cells(f"A{row}:C{row}")
    ws[f"A{row}"].value = v.approver.full_name or (v.approver.username if v.approver else "")
    ws[f"A{row}"].font = font(size=11)
    ws[f"A{row}"].alignment = align("center")
    ws.row_dimensions[row].height = 20

    # ── Xuất stream ─────────────────────────────────────────────────────────
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    fname = f"phieu_{v.code or v.id}_{v.type}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{fname}"',
            "Access-Control-Expose-Headers": "Content-Disposition",
        }
    )